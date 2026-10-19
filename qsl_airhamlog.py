#!/usr/bin/env python3
"""QSLカード レポート面 自動印刷ツール (AirHamlog xlsx -> PDF)

事前印刷済みのQSLカード(100x148mm・縦)に、交信データだけを重ね刷りするPDFを作る。
  - 1交信 = 1ページ(1カード)
  - 印刷済みの重複防止: ログの id を printed_airhamlog.json に記録 (QSL#は本アプリ独自の連番)
  - 海外局(日本のコールサイン以外)は対象外
  - 依存: reportlab, openpyxl (pip install reportlab openpyxl) / GUIは標準のtkinter

使い方:  python qsl_airhamlog.py           -> GUI
         python qsl_airhamlog.py log.xlsx  -> CLI(未印刷分をすべてPDF化)
"""
import json
import re
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas

# ---------------------------------------------------------------- 設定 ----
APP_DIR = Path(__file__).resolve().parent
HISTORY_FILE = APP_DIR / "printed_airhamlog.json"   # 印刷履歴(重複チェック + QSL#採番)
CONFIG_FILE = APP_DIR / "qsl_airhamlog_config.json"

DEFAULT_CONFIG = {
    "rig": "",            # 例: "IC-7300"
    "output_w": "",       # 例: "50"
    "ant": "",            # 自動判定できない周波数(50MHzなど)のとき使うAnt
    "qth": "",            # 例: "JCC3608 三豊(香川)"  ← ログのsent_qthは使わない
    "op_name": "",        # OP.欄の名前
    "rmks_tnx": "TNX FB QSO.",
    "rmks_line2": "次回もよろしくお願いします。",
    "rmks_line3": "",
    "source_tz": "JST",   # ログの時刻 JST / UTC (AirHamlogは「UTC」表記でも実際はJST)
    "print_tz": "JST",    # カードに印字する時刻 JST / UTC
    "template_image": "", # 位置確認用: カード白紙スキャン画像。空なら重ね刷りのみ
    "offset_x_mm": 0.0,   # プリンタのズレ補正
    "offset_y_mm": 0.0,
    "font": "Helvetica",
    "jp_font_path": "",   # 日本語フォント(.ttf/.ttc)。空なら自動検出→無ければ内蔵CIDフォント
}

JP_FONT_CANDIDATES = [
    "C:/Windows/Fonts/meiryo.ttc", "C:/Windows/Fonts/msgothic.ttc",
    "C:/Windows/Fonts/YuGothM.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf",
    "/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
    "/usr/share/fonts/truetype/ipafont-gothic/ipag.ttf",
]

# 白紙スキャン(896x1340px)上の座標。原点=左上。ズレたらここを調整。
SCAN_W, SCAN_H = 896, 1340
PAGE_W_MM, PAGE_H_MM = 100.0, 148.0
POS = {
    "call_box_x": [388, 466, 544, 622, 700, 803],  # TO RADIO 赤枠6つの中心x
    "call_box_y": 135,
    "to_line_x": 62, "to_line_y": 166,              # TO RADIO下の点線欄(移動局: コール/5 など)
    "row_y": 352,
    "date_x": 62, "time_x": 305, "rst_x": 492, "mhz_x": 610, "twoway_x": 785,
    "qslno_x": 160, "qslno_y": 432,
    "rig_x": 155, "rig_y": 492, "out_x": 690,
    "ant_x": 155, "ant_y": 553,
    "rmks_x": 150, "rmks_width": 700,               # Rmks.欄の左端と折り返し幅
    "rmks_y0": 606,                                 # 1行目のベースライン
    "free_x": 75, "free_y1": 730, "free_dy": 48,    # OP. / QTH.
}
# Rmks.の行レイアウト候補: (文字サイズpt, 最大行数, 行間px)。収まらなければ次の候補へ
RMKS_LAYOUTS = [(10, 3, 26), (9, 4, 20)]

MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN",
          "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]

ANT_DP = "1/2λDP"
ANT_HF = "1λワイヤーアンテナ"
ANT_GP = "GPアンテナ"

# 日本のコールサイン: JA-JS / 7J-7N / 8J-8N + 数字 + サフィックス
_JA_CALL = re.compile(r"^(?:J[A-S]|7[J-N]|8[J-N])\d[A-Z0-9]{1,4}$")


# ------------------------------------------------------------ ログ読込 -----
def s(v):
    """セル値を文字列に(Excelの 5 → 5.0 対策)。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_dt(v):
    if isinstance(v, datetime):
        return v
    t = s(v)
    for fmt in ("%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(t, fmt)
        except ValueError:
            pass
    return None


def parse_freq(v):
    """'18MHz' / '3.5MHz' / 18 -> 18.0 (MHz)"""
    m = re.search(r"\d+(?:\.\d+)?", s(v))
    return float(m.group()) if m else None


def normalize(rec):
    mode = re.sub(r"\(.*", "", s(rec.get("mode"))).strip().upper()   # SSB(USB) -> SSB
    return {
        "id": s(rec.get("id")),
        "call": s(rec.get("callsign")).upper(),
        "portable": s(rec.get("portable")).lstrip("/"),               # 「5」→ /5
        "dt": parse_dt(rec.get("qso_at")),
        "rst": s(rec.get("sent_rst")),
        "freq": parse_freq(rec.get("frequency")),
        "mode": mode,
        "qra": unicodedata.normalize("NFKC", s(rec.get("received_qra"))).replace(" ", ""),
        "card_remarks": s(rec.get("card_remarks")),
        # is_qsl_sent / remarks / received_qth / sent_qth は使わない
    }


def read_text(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp932"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def read_log(path):
    """AirHamlogのxlsx(または文字化けしていないCSV)を読み、qso_at昇順の交信リストを返す。"""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    else:
        import csv
        rows = list(csv.reader(read_text(p).splitlines()))
    if not rows:
        return []
    header = [s(h).lower() for h in rows[0]]
    qsos = []
    for r in rows[1:]:
        rec = {h: v for h, v in zip(header, r) if h}
        q = normalize(rec)
        if q["call"] and q["dt"]:
            qsos.append(q)
    qsos.sort(key=lambda q: q["dt"])
    return qsos


def is_japanese(call):
    base = split_call(call)[0]
    return bool(_JA_CALL.match(base))


def split_call(call):
    """'JA1ZGP/5' -> ('JA1ZGP', '5')。"""
    parts = [p for p in call.split("/") if p]
    if len(parts) <= 1:
        return call, ""
    base = max(parts, key=len)
    return base, "/".join(p for p in parts if p is not base)


def qso_key(q):
    """重複判定キー: ログの id(無ければ コール+日時+周波数+モード)"""
    if q["id"]:
        return f"id:{q['id']}"
    return f"{q['call']}|{q['dt']:%Y%m%d%H%M}|{q['freq']}|{q['mode']}"


# ------------------------------------------------------------ 変換 ---------
def fmt_datetime(q, source_tz, print_tz):
    dt = q["dt"]
    if source_tz != print_tz:
        dt += timedelta(hours=-9 if source_tz == "JST" else 9)
    return f"{dt.day:02d} {MONTHS[dt.month - 1]} {dt.year}", dt.strftime("%H%M"), print_tz


def mhz_text(q):
    return f"{q['freq']:g}" if q["freq"] is not None else ""


def ant_of(q, cfg):
    f = q["freq"]
    if f is None:
        return cfg["ant"]
    if 3.4 <= f <= 4.0:
        return ANT_DP          # 3.5MHz
    if f < 30:
        return ANT_HF          # その他のHF
    if 144 <= f <= 146 or 430 <= f <= 440:
        return ANT_GP          # 2m / 430
    return cfg["ant"]


# ----------------------------------------------------------- 履歴/設定 -----
def load_json(path, default):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def load_history():
    return load_json(HISTORY_FILE, {"next_no": 1, "printed": {}})


def save_history(h):
    HISTORY_FILE.write_text(json.dumps(h, ensure_ascii=False, indent=1), encoding="utf-8")


def load_config():
    c = dict(DEFAULT_CONFIG)
    c.update(load_json(CONFIG_FILE, {}))
    return c


def save_config(c):
    CONFIG_FILE.write_text(json.dumps(c, ensure_ascii=False, indent=1), encoding="utf-8")


# ------------------------------------------------------------ フォント -----
_JP_CACHE = {}


def register_jp_font(cfg):
    path = cfg.get("jp_font_path", "")
    if path in _JP_CACHE:
        return _JP_CACHE[path]
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    name = None
    for p in ([path] if path else []) + JP_FONT_CANDIDATES:
        if p and Path(p).exists():
            try:
                pdfmetrics.registerFont(TTFont("QSL-JP", p, subfontIndex=0))
                name = "QSL-JP"
                break
            except Exception:
                continue
    if name is None:
        try:
            pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
            name = "HeiseiKakuGo-W5"
        except Exception:
            name = None
    _JP_CACHE[path] = name
    return name


# ------------------------------------------------------------ Rmks. --------
def wrap_text(tokens, max_w, measure):
    """tokens(分割不可の単位)を max_w に収まるよう貪欲に折り返す。"""
    lines, cur = [], ""
    for t in tokens:
        if cur and measure(cur + t) > max_w:
            lines.append(cur.rstrip())
            cur = t.lstrip()
        else:
            cur += t
    if cur.strip():
        lines.append(cur.rstrip())
    return lines


def rmks_layout(q, cfg, measure_for, max_w):
    """Rmks.欄の行リストを作る。
    先頭ブロック: 「OP名さん+card_remarks TNX FB QSO.」(長ければ折り返し)
    続き: 設定のrmks_line2 / rmks_line3(空なら無し)
    返り値: (行リスト, 文字サイズpt, 行間px)"""
    head = f"{q['qra']}さん" if q["qra"] else ""
    if q["card_remarks"]:
        head += (" " if head else "") + q["card_remarks"]
    tokens = list(head) + ([" "] if head else []) + [cfg["rmks_tnx"]]
    fixed = [t for t in (cfg["rmks_line2"], cfg["rmks_line3"]) if t.strip()]
    for size, max_lines, pitch in RMKS_LAYOUTS:
        lines = wrap_text(tokens, max_w, lambda x: measure_for(x, size)) + fixed
        if len(lines) <= max_lines:
            return lines, size, pitch
    size, max_lines, pitch = RMKS_LAYOUTS[-1]
    return lines[:max_lines], size, pitch


# ------------------------------------------------------------- PDF ---------
def render_pdf(qsos, out_path, cfg, history, assign_numbers=True):
    c = canvas.Canvas(str(out_path), pagesize=(PAGE_W_MM * mm, PAGE_H_MM * mm))
    sx, sy = PAGE_W_MM * mm / SCAN_W, PAGE_H_MM * mm / SCAN_H
    ox, oy = cfg["offset_x_mm"] * mm, -cfg["offset_y_mm"] * mm
    font = cfg["font"]
    jp = register_jp_font(cfg)

    def font_for(sv):
        return jp if (jp and not sv.isascii()) else font

    def X(px): return px * sx + ox
    def Y(py): return (SCAN_H - py) * sy + oy

    def text(px, py, sv, size=11, center=False):
        c.setFont(font_for(sv), size)
        (c.drawCentredString if center else c.drawString)(X(px), Y(py), sv)

    def measure_for(sv, size):
        return pdfmetrics.stringWidth(sv, font_for(sv), size)

    for q in qsos:
        if cfg["template_image"] and Path(cfg["template_image"]).exists():
            c.drawImage(cfg["template_image"], 0, 0, PAGE_W_MM * mm, PAGE_H_MM * mm)
        c.setFillColorRGB(0, 0, 0)

        # TO RADIO: 赤枠は本体のみ、移動局は点線欄に「コール/5」
        base, suffix = split_call(q["call"])
        suffix = q["portable"] or suffix
        for i, ch in enumerate(base[:6]):
            text(POS["call_box_x"][i], POS["call_box_y"] + 14, ch, 20, center=True)
        if suffix:
            text(POS["to_line_x"], POS["to_line_y"], f"{base}/{suffix}", 14)

        date, hhmm, tz = fmt_datetime(q, cfg["source_tz"], cfg["print_tz"])
        text(POS["date_x"], POS["row_y"], date, 11)
        text(POS["time_x"], POS["row_y"], f"{hhmm} {tz}", 11)
        text(POS["rst_x"], POS["row_y"], q["rst"], 12)
        text(POS["mhz_x"], POS["row_y"], mhz_text(q), 11)
        text(POS["twoway_x"], POS["row_y"], q["mode"], 11, center=True)

        key = qso_key(q)
        rec = history["printed"].get(key)
        if rec is None and assign_numbers:
            rec = {"qsl_no": history["next_no"], "call": q["call"], "log_id": q["id"],
                   "printed_at": datetime.now().isoformat(timespec="seconds")}
            history["next_no"] += 1
            history["printed"][key] = rec
        if rec:
            text(POS["qslno_x"], POS["qslno_y"], f"{rec['qsl_no']:05d}", 11)

        text(POS["rig_x"], POS["rig_y"], cfg["rig"], 11)
        text(POS["out_x"], POS["rig_y"], cfg["output_w"], 11)
        text(POS["ant_x"], POS["ant_y"], ant_of(q, cfg), 11)

        max_w = POS["rmks_width"] * sx
        lines, size, pitch = rmks_layout(q, cfg, measure_for, max_w)
        for i, line in enumerate(lines):
            text(POS["rmks_x"], POS["rmks_y0"] + i * pitch, line, size)

        text(POS["free_x"], POS["free_y1"], f"OP.  {cfg['op_name']}", 11)
        text(POS["free_x"], POS["free_y1"] + POS["free_dy"], f"QTH. {cfg['qth']}", 11)
        c.showPage()
    c.save()
    return history


# ------------------------------------------------------------- CLI ---------
def run_cli(path):
    cfg, hist = load_config(), load_history()
    allq = read_log(path)
    foreign = [q for q in allq if not is_japanese(q["call"])]
    todo = [q for q in allq if is_japanese(q["call"]) and qso_key(q) not in hist["printed"]]
    print(f"全{len(allq)}件 / 海外局(対象外){len(foreign)}件 / 未印刷 {len(todo)}件")
    if not todo:
        return
    out = APP_DIR / f"qsl_{datetime.now():%Y%m%d_%H%M%S}.pdf"
    save_history(render_pdf(todo, out, cfg, hist))
    print("出力:", out)


# ------------------------------------------------------------- GUI ---------
def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    cfg = load_config()
    root = tk.Tk()
    root.title("QSLカード レポート面 印刷ツール (AirHamlog)")
    state = {"qsos": []}

    top = ttk.Frame(root, padding=6)
    top.pack(fill="x")
    ttk.Button(top, text="xlsxを開く", command=lambda: load()).pack(side="left")
    ttk.Button(top, text="設定", command=lambda: settings()).pack(side="left", padx=6)
    ttk.Button(top, text="PDF作成", command=lambda: make()).pack(side="right")
    info = tk.StringVar()
    ttk.Label(top, textvariable=info).pack(side="left", padx=10)

    cols = ("state", "id", "call", "date", "time", "mhz", "mode", "rst", "name")
    tree = ttk.Treeview(root, columns=cols, show="headings", selectmode="extended", height=18)
    for col, w, label in zip(cols, (70, 90, 110, 100, 70, 50, 60, 40, 110),
                             ("状態", "ログID", "CALL", "日付", "時刻", "MHz", "MODE", "RST", "相手名")):
        tree.heading(col, text=label)
        tree.column(col, width=w, anchor="w")
    tree.pack(fill="both", expand=True, padx=6, pady=(0, 6))

    def load():
        p = filedialog.askopenfilename(filetypes=[("AirHamlog xlsx", "*.xlsx"), ("CSV", "*.csv"), ("all", "*.*")])
        if not p:
            return
        try:
            state["qsos"] = read_log(p)
        except Exception as e:
            return messagebox.showerror("読み込みエラー", str(e))
        refresh()

    def refresh():
        hist = load_history()
        tree.delete(*tree.get_children())
        todo, nforeign = [], 0
        for i, q in enumerate(state["qsos"]):
            if not is_japanese(q["call"]):
                st, nforeign = "対象外", nforeign + 1
            elif qso_key(q) in hist["printed"]:
                st = "印刷済"
            else:
                st = "未印刷"
                todo.append(str(i))
            d, t, tz = fmt_datetime(q, cfg["source_tz"], cfg["print_tz"])
            tree.insert("", "end", iid=str(i), values=(
                st, q["id"], q["call"], d, f"{t} {tz}", mhz_text(q), q["mode"], q["rst"], q["qra"]))
        tree.selection_set(todo)  # 未印刷だけ選択(=重複防止)
        info.set(f"{len(state['qsos'])}件中 未印刷 {len(todo)}件 / 海外局(対象外) {nforeign}件")

    def make():
        sel = [state["qsos"][int(i)] for i in tree.selection()]
        sel = [q for q in sel if is_japanese(q["call"])]
        if not sel:
            return messagebox.showinfo("QSL", "対象がありません")
        hist = load_history()
        again = [q for q in sel if qso_key(q) in hist["printed"]]
        if again and not messagebox.askyesno("再印刷", f"印刷済みが{len(again)}件含まれます。再印刷しますか?"):
            return
        out = filedialog.asksaveasfilename(defaultextension=".pdf", initialfile=f"qsl_{datetime.now():%Y%m%d_%H%M}.pdf")
        if not out:
            return
        save_history(render_pdf(sel, out, cfg, hist))
        refresh()
        messagebox.showinfo("QSL", f"{len(sel)}枚分のPDFを作成しました")

    def settings():
        win = tk.Toplevel(root)
        win.title("設定")
        labels = [("rig", "Rig"), ("output_w", "Output (W)"),
                  ("ant", "Ant(自動判定外の周波数用)"), ("qth", "QTH"), ("op_name", "OP名"),
                  ("rmks_tnx", "Rmks 定型(名前の後ろ)"), ("rmks_line2", "Rmks 2行目"),
                  ("rmks_line3", "Rmks 3行目"),
                  ("source_tz", "ログの時刻 (JST/UTC)"), ("print_tz", "印字する時刻 (JST/UTC)"),
                  ("template_image", "位置確認用の台紙画像"),
                  ("jp_font_path", "日本語フォント(.ttf/.ttc 空=自動)"),
                  ("offset_x_mm", "X補正(mm)"), ("offset_y_mm", "Y補正(mm)")]
        vars_ = {}
        for r, (k, lab) in enumerate(labels):
            ttk.Label(win, text=lab).grid(row=r, column=0, sticky="w", padx=6, pady=2)
            vars_[k] = tk.StringVar(value=str(cfg[k]))
            ttk.Entry(win, textvariable=vars_[k], width=40).grid(row=r, column=1, padx=6)

        def ok():
            for k, v in vars_.items():
                cfg[k] = float(v.get() or 0) if k.startswith("offset") else v.get().strip()
            save_config(cfg)
            win.destroy()
            refresh()
        ttk.Button(win, text="保存", command=ok).grid(row=len(labels), column=1, sticky="e", pady=6)

    root.mainloop()


if __name__ == "__main__":
    run_cli(sys.argv[1]) if len(sys.argv) > 1 else run_gui()
